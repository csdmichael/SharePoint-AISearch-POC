"""Validate count, provenance, and rich structure of the generated Office corpus."""

from __future__ import annotations

import argparse
import json
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from generate_corpus import (
    FRACTION_PERCENT_FIELDS,
    PERCENTAGE_POINT_FIELDS,
    aggregate_for_table,
    aggregate_kpis,
    columns_for_table,
    data_fields,
    format_value,
    group_rows,
    humanize,
    profile_rows,
    sample_rows,
)


FORMAT_ORDER = ("docx", "pptx", "xlsx")


def expected_formats(document_count: int) -> Counter[str]:
    base, remainder = divmod(document_count, len(FORMAT_ORDER))
    return Counter(
        {
            extension: base + (1 if index < remainder else 0)
            for index, extension in enumerate(FORMAT_ORDER)
        }
    )


def table_values(table) -> list[list[str]]:
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def find_table(tables, header: list[str]):
    for table in tables:
        values = table_values(table)
        if values and values[0] == header:
            return values
    raise AssertionError(f"Expected table header not found: {header}")


def expected_schema_rows(columns: list[dict]) -> list[list[str]]:
    return [
        [column["name"], column["type"], str(column["nullable"]), column["comment"]]
        for column in columns
    ]


def validate_docx(
    path: Path,
    source_table: str,
    expected_facts: list[tuple[str, str]],
    expected_profile: list[tuple[str, str, str, str]],
    expected_schema: list[list[str]],
    expected_rows: list[dict],
) -> None:
    document = Document(path)
    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    assert source_table in text, f"{path}: source table missing from text"
    assert "What the data says" in text, f"{path}: decision notes missing"
    assert "Recommended follow-up" in text, f"{path}: recommended follow-up missing"
    assert "Source schema" in text, f"{path}: source schema missing"
    assert len(document.tables) >= 5, f"{path}: expected KPI, profile, data, provenance, and schema tables"
    assert len(document.inline_shapes) >= 2, f"{path}: expected an embedded chart and lineage diagram"
    kpi_table = find_table(document.tables, ["Indicator", "Value"])
    assert kpi_table[1:] == [list(fact) for fact in expected_facts], f"{path}: KPI rows do not match"
    profile_table = find_table(document.tables, ["Measure", "Minimum", "Average / Count", "Maximum"])
    assert profile_table[1:] == [list(row) for row in expected_profile], f"{path}: profile rows do not match"
    schema_table = find_table(document.tables, ["Column", "Type", "Nullable", "Definition"])
    assert schema_table[1:] == expected_schema, f"{path}: schema rows do not match"
    fields = data_fields(expected_rows)[:7]
    source_table = find_table(document.tables, [humanize(field) for field in fields])
    expected_source_rows = [
        [format_value(field, row[field]) for field in fields] for row in expected_rows
    ]
    assert source_table[1:] == expected_source_rows, f"{path}: source rows do not match"
    body_xml = document.element.body.xml
    assert body_xml.count('w:type="page"') >= 3, f"{path}: expected at least four pages"


def validate_pptx(
    path: Path,
    source_table: str,
    expected_facts: list[tuple[str, str]],
    expected_profile: list[tuple[str, str, str, str]],
    expected_schema: list[list[str]],
    expected_rows: list[dict],
) -> None:
    presentation = Presentation(path)
    assert len(presentation.slides) >= 8, f"{path}: expected at least eight slides"
    text_parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text_parts.append(shape.text)
            if shape.has_table:
                text_parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
    text = "\n".join(text_parts)
    assert source_table in text, f"{path}: source table missing from text"
    assert "What the data says" in text, f"{path}: decision slide missing"
    assert "Recommended follow-up" in text, f"{path}: recommended follow-up missing"
    assert "Source schema" in text, f"{path}: source schema slide missing"
    assert any(shape.has_table for slide in presentation.slides for shape in slide.shapes), f"{path}: no table"
    assert any(shape.has_chart for slide in presentation.slides for shape in slide.shapes), f"{path}: no chart"
    indicator_slide = next(
        slide for slide in presentation.slides if any(getattr(shape, "text", "") == "Key indicators" for shape in slide.shapes)
    )
    for label, value in expected_facts:
        assert any(
            label in shape.text.splitlines() and value in shape.text.splitlines()
            for shape in indicator_slide.shapes
            if hasattr(shape, "text")
        ), f"{path}: KPI pair does not match {label}={value}"
    profile_tables = [
        table_values(shape.table)
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_table and table_values(shape.table)[0] == ["Measure", "Minimum", "Average / Count", "Maximum"]
    ]
    assert len(profile_tables) == 1 and profile_tables[0][1:] == [list(row) for row in expected_profile], (
        f"{path}: profile rows do not match"
    )
    schema_rows = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_table:
                values = table_values(shape.table)
                if values[0] == ["Column", "Type", "Nullable", "Definition"]:
                    schema_rows.extend(values[1:])
    assert schema_rows == expected_schema, f"{path}: paginated schema rows do not match"
    fields = data_fields(expected_rows)[:6]
    source_tables = [
        table_values(shape.table)
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_table and table_values(shape.table)[0] == [humanize(field) for field in fields]
    ]
    expected_source_rows = [
        [format_value(field, row[field]) for field in fields] for row in expected_rows[:7]
    ]
    assert len(source_tables) == 1 and source_tables[0][1:] == expected_source_rows, (
        f"{path}: source rows do not match"
    )


def validate_xlsx(
    path: Path,
    source_table: str,
    expected_facts: list[tuple[str, str]],
    expected_profile: list[tuple[str, str, str, str]],
    expected_schema: list[list[str]],
    expected_rows: list[dict],
) -> None:
    workbook = load_workbook(path, read_only=False, data_only=False)
    required_sheets = {
        "Summary",
        "Source Data",
        "Full Profile",
        "Analysis",
        "Data Dictionary",
        "Decision Notes",
        "Lineage",
    }
    assert required_sheets.issubset(workbook.sheetnames), f"{path}: expected seven rich worksheets"
    assert workbook["Summary"]["B4"].value == source_table, f"{path}: source table mismatch"
    assert sum(len(worksheet.tables) for worksheet in workbook.worksheets) >= 1, f"{path}: no Excel table"
    assert sum(len(worksheet._charts) for worksheet in workbook.worksheets) >= 1, f"{path}: no chart"
    assert len(workbook["Analysis"].conditional_formatting) >= 1, f"{path}: no conditional formatting"
    assert workbook["Decision Notes"]["A2"].value, f"{path}: no decision notes"
    assert "Databricks" in workbook["Lineage"]["B3"].value, f"{path}: no lineage diagram"
    summary_pairs = [
        [str(workbook["Summary"].cell(row, 1).value), str(workbook["Summary"].cell(row, 2).value)]
        for row in range(12, 12 + len(expected_facts))
    ]
    assert summary_pairs == [list(fact) for fact in expected_facts], f"{path}: KPI rows do not match"
    full_profile_rows = [
        [str(cell.value) for cell in row]
        for row in workbook["Full Profile"].iter_rows(min_row=2, max_row=1 + len(expected_profile), max_col=4)
    ]
    assert full_profile_rows == [list(row) for row in expected_profile], f"{path}: profile rows do not match"
    dictionary_rows = [
        [str(cell.value) for cell in row]
        for row in workbook["Data Dictionary"].iter_rows(min_row=2, max_col=3)
    ]
    assert dictionary_rows == [[row[0], row[1], row[3]] for row in expected_schema], (
        f"{path}: schema rows do not match"
    )
    source_sheet = workbook["Source Data"]
    fields = data_fields(expected_rows)
    assert [cell.value for cell in source_sheet[1]] == fields, f"{path}: source headers do not match"
    actual_source_rows = [
        [cell.value for cell in row]
        for row in source_sheet.iter_rows(min_row=2, max_row=1 + len(expected_rows), max_col=len(fields))
    ]
    expected_source_rows = [[row[field] for field in fields] for row in expected_rows]
    assert actual_source_rows == expected_source_rows, f"{path}: source rows do not match"
    headers = {cell.value: cell.column for cell in source_sheet[1]}
    for field in FRACTION_PERCENT_FIELDS.intersection(headers):
        assert all(cell.number_format == "0.00%" for cell in list(source_sheet.iter_cols(min_col=headers[field], max_col=headers[field], min_row=2))[0]), (
            f"{path}: fractional percentage format missing for {field}"
        )
    for field in PERCENTAGE_POINT_FIELDS.intersection(headers):
        assert all(cell.number_format == '0.00"%"' for cell in list(source_sheet.iter_cols(min_col=headers[field], max_col=headers[field], min_row=2))[0]), (
            f"{path}: percentage-point format missing for {field}"
        )
    workbook.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate the generated Office corpus.")
    parser.add_argument("--corpus", type=Path, default=Path("corpus"))
    parser.add_argument("--expected-count", type=int, default=100)
    parser.add_argument("--profile", type=Path, default=Path("data/semiconductor_profile.json"))
    args = parser.parse_args()
    if args.expected_count < 1:
        parser.error("--expected-count must be at least 1")
    manifest = json.loads((args.corpus / "manifest.json").read_text(encoding="utf-8"))
    profile = json.loads(args.profile.read_text(encoding="utf-8"))
    grouped_rows = group_rows(profile)
    assert len(manifest) == args.expected_count, (
        f"Expected {args.expected_count} manifest entries, found {len(manifest)}"
    )
    assert len({item["artifact_id"] for item in manifest}) == args.expected_count, "Artifact IDs are not unique"
    assert len({item["relative_path"] for item in manifest}) == args.expected_count, "Paths are not unique"
    expected = expected_formats(args.expected_count)
    assert Counter(item["extension"] for item in manifest) == expected

    files = [
        path
        for path in args.corpus.rglob("*")
        if path.is_file() and path.suffix.lower() in {".docx", ".pptx", ".xlsx"}
    ]
    assert len(files) == args.expected_count, (
        f"Expected {args.expected_count} Office files, found {len(files)}"
    )
    counts: Counter[str] = Counter()
    for path in files:
        relative = path.relative_to(args.corpus).as_posix()
        item = next(entry for entry in manifest if entry["relative_path"] == relative)
        aggregate = aggregate_for_table(profile, item["source_table"])
        columns = columns_for_table(profile, item["source_table"])
        expected_facts = aggregate_kpis(aggregate, [])
        expected_profile = profile_rows(aggregate)
        expected_schema = expected_schema_rows(columns)
        artifact_number = int(item["artifact_id"].split("-")[1])
        expected_rows = sample_rows(grouped_rows[item["source_table"]], artifact_number)
        assert item["full_source_rows"] == aggregate.get("row_count"), (
            f"{path}: manifest row count does not match the source profile"
        )
        assert item["source_column_count"] == len(columns), (
            f"{path}: manifest schema count does not match the source profile"
        )
        assert zipfile.is_zipfile(path), f"{path}: invalid Open XML package"
        assert path.stat().st_size > 5_000, f"{path}: unexpectedly small"
        extension = path.suffix[1:].lower()
        counts[extension] += 1
        if extension == "docx":
            validate_docx(path, item["source_table"], expected_facts, expected_profile, expected_schema, expected_rows)
        elif extension == "pptx":
            validate_pptx(path, item["source_table"], expected_facts, expected_profile, expected_schema, expected_rows)
        else:
            validate_xlsx(path, item["source_table"], expected_facts, expected_profile, expected_schema, expected_rows)
    assert counts == expected
    print(json.dumps({"status": "passed", "documents": len(files), "formats": dict(counts)}))


if __name__ == "__main__":
    main()