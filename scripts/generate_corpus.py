"""Generate 100 rich Office documents from a validated Databricks profile."""

from __future__ import annotations

import argparse
import json
import re
import shutil
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from statistics import mean
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches as DocxInches
from docx.shared import Pt as DocxPt
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

try:
    from .fetch_semiconductor_profile import validate_profile
except ImportError:
    from fetch_semiconductor_profile import validate_profile


AUTHOR = "Michael Yaacoub @ Microsoft"
SOURCE_SYSTEM = "Azure Databricks via Microsoft Foundry Databricks MCP"
FORMAT_ORDER = ("docx", "pptx", "xlsx")
PERSPECTIVES = (
    "Executive Brief",
    "Operational Review",
    "Trend Analysis",
    "Risk Snapshot",
    "Performance Deep Dive",
    "Planning Workbook",
    "Quality Review",
    "Regional Analysis",
    "Management Summary",
)

TABLE_CONFIG = {
    "defect_analysis": ("Quality", "Defect Analysis", "C0392B"),
    "fab_production": ("Manufacturing", "Fab Production", "1565C0"),
    "inventory": ("Inventory", "Inventory Position", "2E7D32"),
    "product_sales": ("Sales", "Product Sales", "B35A00"),
    "supply_chain": ("Supply Chain", "Supply Chain", "6C3483"),
    "wafer_yield": ("Yield", "Wafer Yield", "00796B"),
}

FRACTION_PERCENT_FIELDS = {
    "yield_pct",
    "gross_margin_pct",
    "avg_yield_pct",
    "best_yield_pct",
    "worst_yield_pct",
    "target_yield_pct",
    "yield_vs_target",
}
PERCENTAGE_POINT_FIELDS = {"on_time_delivery_pct"}


def suffix(table_name: str) -> str:
    return table_name.rsplit(".", 1)[-1]


def slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def humanize(value: str) -> str:
    return value.replace("_", " ").title()


def format_value(field: str, value: Any) -> str:
    if value is None:
        return "Not available"
    if isinstance(value, float):
        if field in FRACTION_PERCENT_FIELDS:
            return f"{value:.2%}"
        if field in PERCENTAGE_POINT_FIELDS:
            return f"{value:.2f}%"
        if field.endswith("_usd"):
            return f"${value:,.2f}"
        return f"{value:,.2f}"
    if isinstance(value, int):
        return f"{value:,}"
    return str(value)


def group_rows(profile: dict) -> dict[str, list[dict]]:
    representative = profile["representative_rows"]
    if isinstance(representative.get("rows"), list):
        rows = representative["rows"]
    else:
        rows = [
            row
            for sample in representative.values()
            if isinstance(sample, dict)
            for row in sample.get("rows", [])
        ]
    grouped: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        grouped[row["source_table"]].append(row)
    return dict(grouped)


def sample_rows(rows: list[dict], document_number: int) -> list[dict]:
    sample_size = min(8, len(rows))
    start = (document_number * 3) % len(rows)
    return [rows[(start + offset) % len(rows)] for offset in range(sample_size)]


def data_fields(rows: list[dict]) -> list[str]:
    return [field for field in rows[0] if field != "source_table"]


def numeric_fields(rows: list[dict]) -> list[str]:
    return [
        field
        for field in data_fields(rows)
        if all(isinstance(row.get(field), (int, float)) for row in rows)
    ]


def dimension_fields(rows: list[dict]) -> list[str]:
    return [
        field
        for field in data_fields(rows)
        if any(isinstance(row.get(field), str) for row in rows)
    ]


def aggregate_for_table(profile: dict, table_name: str) -> dict:
    aggregate_root = profile.get("aggregate_metrics", {})
    results = aggregate_root.get("results") if isinstance(aggregate_root, dict) else None
    if isinstance(results, list):
        return next(
            (result for result in results if result.get("source_table") == table_name),
            {},
        )

    candidates = (table_name, ".".join(table_name.split(".")[-2:]), suffix(table_name))
    entry = next(
        (
            aggregate_root[candidate]
            for candidate in candidates
            if isinstance(aggregate_root, dict) and candidate in aggregate_root
        ),
        {},
    )
    values = entry.get("values", entry) if isinstance(entry, dict) else {}
    if "metric_1" in values:
        return {"source_table": table_name, **values}

    metric_groups: dict[str, dict[str, Any]] = defaultdict(dict)
    for key, value in values.items():
        match = re.match(r"^(min|max|avg|mean)_(.+)$", key)
        if match and not match.group(2).endswith(("date", "month")):
            statistic = "avg" if match.group(1) == "mean" else match.group(1)
            metric_groups[match.group(2)][statistic] = value
    complete_metrics = [
        {"column": field, **statistics}
        for field, statistics in metric_groups.items()
        if {"min", "max", "avg"}.issubset(statistics)
    ]
    date_min = next(
        (value for key, value in values.items() if key.startswith("min_") and key.endswith(("date", "month"))),
        None,
    )
    date_max = next(
        (value for key, value in values.items() if key.startswith("max_") and key.endswith(("date", "month"))),
        None,
    )
    distinct_values = [
        {"column": key.removeprefix("distinct_"), "value": value}
        for key, value in values.items()
        if key.startswith("distinct_")
    ]
    return {
        "source_table": table_name,
        "row_count": values.get("row_count"),
        "min_date": date_min,
        "max_date": date_max,
        "null_cell_count": sum(values.get("null_counts", {}).values()),
        "distinct_entity_1": distinct_values[0] if distinct_values else None,
        "distinct_entity_2": distinct_values[1] if len(distinct_values) > 1 else None,
        "metric_1": complete_metrics[0] if complete_metrics else None,
        "metric_2": complete_metrics[1] if len(complete_metrics) > 1 else None,
    }


def columns_for_table(profile: dict, table_name: str) -> list[dict]:
    column_root = profile.get("column_definitions", {})
    tables = column_root.get("tables", column_root) if isinstance(column_root, dict) else {}
    candidates = (table_name, ".".join(table_name.split(".")[-2:]), suffix(table_name))
    raw_columns = next(
        (tables[candidate] for candidate in candidates if isinstance(tables, dict) and candidate in tables),
        [],
    )
    return [
        {
            "name": column.get("column_name", column.get("name", "")),
            "type": column.get("data_type", column.get("type", "unknown")),
            "nullable": column.get("nullable", "unknown"),
            "comment": column.get("comment") or "Source field from the Databricks table.",
        }
        for column in raw_columns
    ]


def aggregate_kpis(aggregate: dict, rows: list[dict]) -> list[tuple[str, str]]:
    result = []
    if aggregate.get("row_count") is not None:
        result.append(("Full Source Rows", format_value("row_count", aggregate["row_count"])))
    if aggregate.get("min_date") and aggregate.get("max_date"):
        result.append(("Coverage", f"{aggregate['min_date']} to {aggregate['max_date']}"))
    for metric_key in ("metric_1", "metric_2"):
        metric = aggregate.get(metric_key)
        if metric:
            field = metric["column"]
            result.append((f"Average {humanize(field)}", format_value(field, metric.get("avg"))))
    if not result:
        result.append(("Representative Rows", str(len(rows))))
    return result[:4]


def profile_rows(aggregate: dict) -> list[tuple[str, str, str, str]]:
    rows = []
    for metric_key in ("metric_1", "metric_2"):
        metric = aggregate.get(metric_key)
        if metric:
            field = metric["column"]
            rows.append(
                (
                    humanize(field),
                    format_value(field, metric.get("min")),
                    format_value(field, metric.get("avg")),
                    format_value(field, metric.get("max")),
                )
            )
    for entity_key in ("distinct_entity_1", "distinct_entity_2"):
        entity = aggregate.get(entity_key)
        if entity:
            rows.append((f"Distinct {humanize(entity['column'])}", "-", str(entity["value"]), "-"))
    return rows


FOLLOW_UPS = {
    "Quality": "Prioritize the highest-PPM and high-severity combinations, then compare them by fab and process node.",
    "Manufacturing": "Compare yield and cycle time by fab, process node, and product family before adjusting production plans.",
    "Inventory": "Review low-stock and below-reorder positions by region, then reconcile reservations against days of supply.",
    "Sales": "Segment revenue, unit volume, price, and margin by region and customer segment before changing commercial priorities.",
    "Supply Chain": "Escalate long-lead or high-risk component suppliers and validate delivery and quality tradeoffs.",
    "Yield": "Investigate below-target yield combinations by fab, process node, and product family before changing process controls.",
}


def decision_notes(spec: dict, rows: list[dict], aggregate: dict) -> list[tuple[str, str]]:
    notes = []
    row_count = aggregate.get("row_count")
    if row_count is not None:
        coverage = ""
        if aggregate.get("min_date") and aggregate.get("max_date"):
            coverage = f" covering {aggregate['min_date']} through {aggregate['max_date']}"
        notes.append(("Full-table fact", f"The source profile contains {row_count:,} records{coverage}."))
    if aggregate.get("null_cell_count") == 0:
        notes.append(("Data quality fact", "The profiling query found zero null cells across inspected source columns."))
    metric = aggregate.get("metric_2") or aggregate.get("metric_1")
    if metric:
        field = metric["column"]
        notes.append(
            (
                "Full-table range",
                f"{humanize(field)} ranges from {format_value(field, metric.get('min'))} to "
                f"{format_value(field, metric.get('max'))}, with an average of "
                f"{format_value(field, metric.get('avg'))}.",
            )
        )
    sample_metric = next((field for field in numeric_fields(rows) if rows and field in rows[0]), None)
    if sample_metric:
        highest = max(rows, key=lambda row: float(row[sample_metric]))
        dimensions = dimension_fields(rows)
        label = str(highest[dimensions[0]]) if dimensions else "the sampled record"
        notes.append(
            (
                "Representative-sample observation",
                f"Within this artifact's bounded sample, {label} has the highest "
                f"{humanize(sample_metric).lower()} at {format_value(sample_metric, highest[sample_metric])}.",
            )
        )
    notes.append(("Recommended follow-up", FOLLOW_UPS[spec["category"]]))
    return notes


def narrative(table_name: str, rows: list[dict], generated_at: str) -> str:
    dimensions = dimension_fields(rows)
    coverage = ""
    if dimensions:
        field = dimensions[0]
        coverage = (
            f" The sample spans {len({row[field] for row in rows})} distinct "
            f"{humanize(field).lower()} values."
        )
    return (
        f"This generated analytical artifact uses {len(rows)} representative records "
        f"from {table_name}.{coverage} Values shown in tables and charts are derived "
        f"from the Databricks profile generated at {generated_at}; this document is not "
        "a replacement for the source-of-record table."
    )


def font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    candidates = (
        Path("C:/Windows/Fonts/aptos.ttf"),
        Path("C:/Windows/Fonts/calibri.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    )
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def create_lineage_image(path: Path, table_name: str, accent: str) -> None:
    image = Image.new("RGB", (1400, 520), "#F4F7F8")
    draw = ImageDraw.Draw(image)
    boxes = (
        (55, 150, 385, 370, "Databricks", table_name),
        (535, 150, 865, 370, "SharePoint", "Semiconductor Knowledge"),
        (1015, 150, 1345, 370, "Azure AI Search", "Chunked hybrid index"),
    )
    for left, top, right, bottom, heading, detail in boxes:
        draw.rounded_rectangle(
            (left, top, right, bottom), radius=18, fill="#FFFFFF", outline=f"#{accent}", width=6
        )
        draw.text((left + 28, top + 35), heading, fill="#15242B", font=font(34, True))
        lines = [detail[index : index + 34] for index in range(0, len(detail), 34)]
        for line_number, line in enumerate(lines[:3]):
            draw.text(
                (left + 28, top + 102 + line_number * 34),
                line,
                fill="#42545C",
                font=font(24),
            )
    for start, end in ((385, 535), (865, 1015)):
        draw.line((start + 25, 260, end - 25, 260), fill=f"#{accent}", width=10)
        draw.polygon(
            ((end - 55, 235), (end - 25, 260), (end - 55, 285)),
            fill=f"#{accent}",
        )
    draw.text((55, 45), "Knowledge ingestion lineage", fill="#15242B", font=font(42, True))
    image.save(path)


def create_metric_chart_image(path: Path, rows: list[dict], spec: dict) -> None:
    candidates = [
        field
        for field in numeric_fields(rows)
        if min(float(row[field]) for row in rows) >= 0
    ]
    metric = candidates[spec["number"] % len(candidates)] if candidates else numeric_fields(rows)[0]
    dimensions = dimension_fields(rows)
    label_field = dimensions[0] if dimensions else data_fields(rows)[0]
    chart_rows = rows[:7]
    values = [float(row[metric]) for row in chart_rows]
    maximum = max(values) or 1
    image = Image.new("RGB", (1400, 720), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.text((55, 35), f"{humanize(metric)} by {humanize(label_field)}", fill="#15242B", font=font(38, True))
    chart_left, chart_top, chart_right, chart_bottom = 110, 120, 1340, 610
    draw.line((chart_left, chart_top, chart_left, chart_bottom), fill="#52646C", width=3)
    draw.line((chart_left, chart_bottom, chart_right, chart_bottom), fill="#52646C", width=3)
    slot_width = (chart_right - chart_left) / len(chart_rows)
    for index, (row, value) in enumerate(zip(chart_rows, values)):
        left = chart_left + index * slot_width + 18
        right = chart_left + (index + 1) * slot_width - 18
        height = (value / maximum) * (chart_bottom - chart_top - 45)
        top = chart_bottom - height
        draw.rounded_rectangle((left, top, right, chart_bottom), radius=8, fill=f"#{spec['accent']}")
        value_text = format_value(metric, row[metric])
        draw.text((left, max(chart_top, top - 30)), value_text, fill="#15242B", font=font(18, True))
        label = str(row[label_field])[:16]
        draw.text((left, chart_bottom + 15), label, fill="#42545C", font=font(17))
    draw.text(
        (55, 680),
        "Representative sample only; full-table metrics are reported separately.",
        fill="#52646C",
        font=font(18),
    )
    image.save(path)


def set_docx_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = DocxPt(9)


def create_docx(
    path: Path,
    spec: dict,
    rows: list[dict],
    aggregate: dict,
    columns: list[dict],
    diagram_path: Path,
    chart_path: Path,
) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = DocxInches(0.65)
    section.bottom_margin = DocxInches(0.65)
    section.left_margin = DocxInches(0.7)
    section.right_margin = DocxInches(0.7)
    document.core_properties.author = AUTHOR
    document.core_properties.title = spec["title"]
    document.core_properties.subject = spec["category"]
    document.core_properties.keywords = f"semiconductor, {spec['category']}, Databricks, Azure AI Search"

    title = document.add_heading(spec["title"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(f"{spec['category']} | Generated knowledge artifact")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f"Prepared by {AUTHOR}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("Executive context", level=1)
    document.add_paragraph(narrative(spec["source_table"], rows, spec["profile_generated_at"]))
    document.add_heading("Purpose", level=1)
    document.add_paragraph(
        f"This {spec['perspective'].lower()} supports decisions in {spec['category'].lower()} by combining "
        "full-table profile metrics with a bounded, traceable sample of source records."
    )
    document.add_heading("Key full-table indicators", level=1)
    indicator_table = document.add_table(rows=1, cols=2)
    indicator_table.style = "Light Shading Accent 1"
    set_docx_cell_text(indicator_table.rows[0].cells[0], "Indicator", True)
    set_docx_cell_text(indicator_table.rows[0].cells[1], "Value", True)
    for name, value in aggregate_kpis(aggregate, rows):
        cells = indicator_table.add_row().cells
        set_docx_cell_text(cells[0], name)
        set_docx_cell_text(cells[1], value)

    document.add_page_break()
    document.add_heading("Source profile and decision notes", level=1)
    profile_table = document.add_table(rows=1, cols=4)
    profile_table.style = "Light Shading Accent 1"
    for column, heading in enumerate(("Measure", "Minimum", "Average / Count", "Maximum")):
        set_docx_cell_text(profile_table.rows[0].cells[column], heading, True)
    for metric_name, minimum, average, maximum in profile_rows(aggregate):
        cells = profile_table.add_row().cells
        for column, value in enumerate((metric_name, minimum, average, maximum)):
            set_docx_cell_text(cells[column], value)
    document.add_picture(str(chart_path), width=DocxInches(7.3))
    document.add_heading("What the data says", level=2)
    for heading, detail in decision_notes(spec, rows, aggregate):
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(f"{heading}: ").bold = True
        paragraph.add_run(detail)

    document.add_page_break()
    document.add_heading("Representative source records", level=1)
    document.add_paragraph(
        "The following bounded sample is included to support retrieval, table extraction, "
        "and source-grounded question answering."
    )
    fields = data_fields(rows)[:7]
    data_table = document.add_table(rows=1, cols=len(fields))
    data_table.style = "Light Shading Accent 1"
    for column, field in enumerate(fields):
        set_docx_cell_text(data_table.rows[0].cells[column], humanize(field), True)
    for row in rows:
        cells = data_table.add_row().cells
        for column, field in enumerate(fields):
            set_docx_cell_text(cells[column], format_value(field, row.get(field)))

    document.add_page_break()
    document.add_heading("Lineage and retrieval design", level=1)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(diagram_path), width=DocxInches(7.3))
    caption = document.add_paragraph(
        "Diagram: Databricks facts are rendered into Office documents, governed in "
        "SharePoint, then chunked and vectorized by the SharePoint indexer."
    )
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_heading("Source provenance", level=2)
    provenance = document.add_table(rows=0, cols=2)
    provenance.style = "Light Shading Accent 1"
    for label, value in (
        ("Source system", SOURCE_SYSTEM),
        ("Source table", spec["source_table"]),
        ("Profile timestamp", spec["profile_generated_at"]),
        ("Artifact ID", spec["artifact_id"]),
    ):
        cells = provenance.add_row().cells
        set_docx_cell_text(cells[0], label, True)
        set_docx_cell_text(cells[1], value)
    document.add_heading("Source schema", level=2)
    schema_table = document.add_table(rows=1, cols=4)
    schema_table.style = "Light Shading Accent 1"
    for column, heading in enumerate(("Column", "Type", "Nullable", "Definition")):
        set_docx_cell_text(schema_table.rows[0].cells[column], heading, True)
    for column_definition in columns:
        cells = schema_table.add_row().cells
        for column, value in enumerate(
            (
                column_definition["name"],
                column_definition["type"],
                str(column_definition["nullable"]),
                column_definition["comment"],
            )
        ):
            set_docx_cell_text(cells[column], value)
    footer = section.footer.paragraphs[0]
    footer.text = f"{AUTHOR} | {spec['source_table']} | {spec['artifact_id']}"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.save(path)


def add_ppt_footer(slide, spec: dict) -> None:
    box = slide.shapes.add_textbox(Inches(0.35), Inches(7.15), Inches(12.6), Inches(0.25))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = f"{spec['source_table']} | {spec['profile_generated_at']} | {spec['artifact_id']}"
    paragraph.font.size = Pt(8)
    paragraph.font.color.rgb = RGBColor(88, 103, 110)


def add_ppt_title(slide, text: str, accent: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12.0), Inches(0.6))
    paragraph = box.text_frame.paragraphs[0]
    paragraph.text = text
    paragraph.font.size = Pt(27)
    paragraph.font.bold = True
    paragraph.font.color.rgb = RGBColor.from_string(accent)


def create_pptx(path: Path, spec: dict, rows: list[dict], aggregate: dict, columns: list[dict]) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.author = AUTHOR
    presentation.core_properties.title = spec["title"]
    blank = presentation.slide_layouts[6]
    accent = spec["accent"]

    slide = presentation.slides.add_slide(blank)
    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, presentation.slide_width, presentation.slide_height)
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor.from_string(accent)
    banner.line.fill.background()
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.65), Inches(11.7), Inches(2.0))
    frame = title_box.text_frame
    frame.text = spec["title"]
    frame.paragraphs[0].font.size = Pt(38)
    frame.paragraphs[0].font.bold = True
    frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    paragraph = frame.add_paragraph()
    paragraph.text = f"{spec['category']} | Source-grounded semiconductor briefing"
    paragraph.font.size = Pt(20)
    paragraph.font.color.rgb = RGBColor(238, 244, 246)
    source = frame.add_paragraph()
    source.text = spec["source_table"]
    source.font.size = Pt(13)
    source.font.color.rgb = RGBColor(238, 244, 246)

    slide = presentation.slides.add_slide(blank)
    add_ppt_title(slide, "Key indicators", accent)
    for index, (name, value) in enumerate(aggregate_kpis(aggregate, rows)):
        left = 0.7 + (index % 2) * 6.15
        top = 1.35 + (index // 2) * 2.35
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(5.55), Inches(1.75))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(246, 249, 250)
        shape.line.color.rgb = RGBColor.from_string(accent)
        frame = shape.text_frame
        frame.text = value
        frame.paragraphs[0].font.size = Pt(27)
        frame.paragraphs[0].font.bold = True
        frame.paragraphs[0].font.color.rgb = RGBColor.from_string(accent)
        label = frame.add_paragraph()
        label.text = name
        label.font.size = Pt(13)
        label.font.color.rgb = RGBColor(64, 80, 88)
    add_ppt_footer(slide, spec)

    slide = presentation.slides.add_slide(blank)
    add_ppt_title(slide, "Full source profile", accent)
    profile_data = profile_rows(aggregate)
    profile_shape = slide.shapes.add_table(
        len(profile_data) + 1,
        4,
        Inches(0.75),
        Inches(1.25),
        Inches(11.8),
        Inches(4.8),
    )
    profile_table = profile_shape.table
    for column, heading in enumerate(("Measure", "Minimum", "Average / Count", "Maximum")):
        profile_table.cell(0, column).text = heading
    for row_number, profile_row in enumerate(profile_data, start=1):
        for column, value in enumerate(profile_row):
            profile_table.cell(row_number, column).text = value
    for row_number in range(len(profile_table.rows)):
        for column in range(len(profile_table.columns)):
            cell = profile_table.cell(row_number, column)
            cell.text_frame.paragraphs[0].font.size = Pt(11)
            if row_number == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(accent)
                cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
                cell.text_frame.paragraphs[0].font.bold = True
    add_ppt_footer(slide, spec)

    slide = presentation.slides.add_slide(blank)
    add_ppt_title(slide, "What the data says", accent)
    for index, (heading, detail) in enumerate(decision_notes(spec, rows, aggregate)):
        top = 1.25 + index * 1.18
        marker = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.75), Inches(top), Inches(0.42), Inches(0.42))
        marker.fill.solid()
        marker.fill.fore_color.rgb = RGBColor.from_string(accent)
        marker.line.fill.background()
        text_box = slide.shapes.add_textbox(Inches(1.35), Inches(top - 0.05), Inches(10.9), Inches(0.9))
        frame = text_box.text_frame
        frame.text = heading
        frame.paragraphs[0].font.size = Pt(16)
        frame.paragraphs[0].font.bold = True
        frame.paragraphs[0].font.color.rgb = RGBColor.from_string(accent)
        detail_paragraph = frame.add_paragraph()
        detail_paragraph.text = detail
        detail_paragraph.font.size = Pt(12)
        detail_paragraph.font.color.rgb = RGBColor(64, 80, 88)
    add_ppt_footer(slide, spec)

    slide = presentation.slides.add_slide(blank)
    add_ppt_title(slide, "Representative metric view", accent)
    numbers = numeric_fields(rows)
    dimensions = dimension_fields(rows)
    if numbers:
        metric = numbers[spec["number"] % len(numbers)]
        label_field = dimensions[0] if dimensions else data_fields(rows)[0]
        chart_data = ChartData()
        chart_data.categories = [str(row[label_field])[:24] for row in rows[:7]]
        chart_data.add_series(humanize(metric), [float(row[metric]) for row in rows[:7]])
        chart = slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(0.8),
            Inches(1.25),
            Inches(11.8),
            Inches(5.25),
            chart_data,
        ).chart
        chart.has_legend = False
        chart.value_axis.has_major_gridlines = True
        chart.chart_title.text_frame.text = f"{humanize(metric)} by {humanize(label_field)}"
    add_ppt_footer(slide, spec)

    schema_page_size = 8
    schema_page_count = max(1, (len(columns) + schema_page_size - 1) // schema_page_size)
    for page_index in range(schema_page_count):
        slide = presentation.slides.add_slide(blank)
        add_ppt_title(slide, f"Source schema ({page_index + 1}/{schema_page_count})", accent)
        visible_columns = columns[page_index * schema_page_size : (page_index + 1) * schema_page_size]
        schema_shape = slide.shapes.add_table(
            len(visible_columns) + 1,
            4,
            Inches(0.7),
            Inches(1.15),
            Inches(11.9),
            Inches(5.65),
        )
        schema_table = schema_shape.table
        for column, heading in enumerate(("Column", "Type", "Nullable", "Definition")):
            schema_table.cell(0, column).text = heading
        for row_number, column_definition in enumerate(visible_columns, start=1):
            for column, value in enumerate(
                (
                    column_definition["name"],
                    column_definition["type"],
                    str(column_definition["nullable"]),
                    column_definition["comment"],
                )
            ):
                schema_table.cell(row_number, column).text = value
        for row_number in range(len(schema_table.rows)):
            for column in range(len(schema_table.columns)):
                cell = schema_table.cell(row_number, column)
                cell.text_frame.paragraphs[0].font.size = Pt(9)
                if row_number == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(accent)
                    cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
                    cell.text_frame.paragraphs[0].font.bold = True
        add_ppt_footer(slide, spec)

    slide = presentation.slides.add_slide(blank)
    add_ppt_title(slide, "Representative records", accent)
    fields = data_fields(rows)[:6]
    table_shape = slide.shapes.add_table(
        min(7, len(rows)) + 1,
        len(fields),
        Inches(0.45),
        Inches(1.15),
        Inches(12.4),
        Inches(5.55),
    )
    table = table_shape.table
    for column, field in enumerate(fields):
        table.cell(0, column).text = humanize(field)
    for row_number, row in enumerate(rows[:7], start=1):
        for column, field in enumerate(fields):
            table.cell(row_number, column).text = format_value(field, row[field])
    for row_number in range(len(table.rows)):
        for column in range(len(table.columns)):
            cell = table.cell(row_number, column)
            cell.text_frame.paragraphs[0].font.size = Pt(9)
            if row_number == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(accent)
                cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
                cell.text_frame.paragraphs[0].font.bold = True
    add_ppt_footer(slide, spec)

    slide = presentation.slides.add_slide(blank)
    add_ppt_title(slide, "Knowledge lineage", accent)
    stages = (
        (0.65, "Databricks", suffix(spec["source_table"])),
        (4.65, "SharePoint", "Semiconductor Knowledge"),
        (8.65, "Azure AI Search", "Chunk + vector + hybrid"),
    )
    for index, (left, heading, detail) in enumerate(stages):
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(2.1), Inches(3.35), Inches(2.15))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(246, 249, 250)
        box.line.color.rgb = RGBColor.from_string(accent)
        frame = box.text_frame
        frame.text = heading
        frame.paragraphs[0].font.size = Pt(19)
        frame.paragraphs[0].font.bold = True
        frame.paragraphs[0].font.color.rgb = RGBColor.from_string(accent)
        detail_paragraph = frame.add_paragraph()
        detail_paragraph.text = detail
        detail_paragraph.font.size = Pt(12)
        detail_paragraph.alignment = PP_ALIGN.CENTER
        if index < 2:
            arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(left + 3.45), Inches(2.72), Inches(0.8), Inches(0.8))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = RGBColor.from_string(accent)
            arrow.line.fill.background()
    note = slide.shapes.add_textbox(Inches(1.0), Inches(5.1), Inches(11.3), Inches(0.9))
    note.text_frame.text = narrative(spec["source_table"], rows, spec["profile_generated_at"])
    note.text_frame.paragraphs[0].font.size = Pt(12)
    note.text_frame.paragraphs[0].font.color.rgb = RGBColor(64, 80, 88)
    add_ppt_footer(slide, spec)
    presentation.save(path)


def style_excel_header(cells, accent: str) -> None:
    for cell in cells:
        cell.fill = PatternFill("solid", fgColor=accent)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def create_xlsx(path: Path, spec: dict, rows: list[dict], aggregate: dict, columns: list[dict]) -> None:
    workbook = Workbook()
    workbook.properties.creator = AUTHOR
    workbook.properties.title = spec["title"]
    workbook.properties.subject = spec["category"]
    workbook.properties.keywords = "semiconductor, Databricks, SharePoint, Azure AI Search"
    accent = spec["accent"]

    summary = workbook.active
    summary.title = "Summary"
    summary.merge_cells("A1:F2")
    summary["A1"] = spec["title"]
    summary["A1"].font = Font(size=20, bold=True, color="FFFFFF")
    summary["A1"].fill = PatternFill("solid", fgColor=accent)
    summary["A1"].alignment = Alignment(vertical="center")
    summary["A4"] = "Source table"
    summary["B4"] = spec["source_table"]
    summary["A5"] = "Profile generated"
    summary["B5"] = spec["profile_generated_at"]
    summary["A6"] = "Artifact ID"
    summary["B6"] = spec["artifact_id"]
    summary["A8"] = "Narrative"
    summary["B8"] = narrative(spec["source_table"], rows, spec["profile_generated_at"])
    summary.merge_cells("B8:F10")
    summary["B8"].alignment = Alignment(wrap_text=True, vertical="top")
    for row_number, (name, value) in enumerate(aggregate_kpis(aggregate, rows), start=12):
        summary.cell(row_number, 1, name)
        summary.cell(row_number, 2, value)
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 32
    for column in "CDEF":
        summary.column_dimensions[column].width = 15

    source = workbook.create_sheet("Source Data")
    fields = data_fields(rows)
    source.append(fields)
    for row in rows:
        source.append([row[field] for field in fields])
    style_excel_header(source[1], accent)
    source.freeze_panes = "A2"
    source.auto_filter.ref = source.dimensions
    for index, field in enumerate(fields, start=1):
        source.column_dimensions[source.cell(1, index).column_letter].width = min(24, max(13, len(field) + 2))
        number_format = None
        if field in FRACTION_PERCENT_FIELDS:
            number_format = "0.00%"
        elif field in PERCENTAGE_POINT_FIELDS:
            number_format = '0.00"%"'
        elif field.endswith("_usd"):
            number_format = '$#,##0.00'
        elif field in numeric_fields(rows):
            number_format = '#,##0.00'
        if number_format:
            for cell in source.iter_cols(min_col=index, max_col=index, min_row=2, max_row=len(rows) + 1):
                for value_cell in cell:
                    value_cell.number_format = number_format
    table = Table(displayName=f"SourceData{spec['number']:03d}", ref=source.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    source.add_table(table)

    full_profile = workbook.create_sheet("Full Profile")
    full_profile.append(["Measure", "Minimum", "Average / Count", "Maximum"])
    for metric_name, minimum, average, maximum in profile_rows(aggregate):
        full_profile.append([metric_name, minimum, average, maximum])
    full_profile.append(["Null cells", "-", aggregate.get("null_cell_count", "Unknown"), "-"])
    style_excel_header(full_profile[1], accent)
    for column in "ABCD":
        full_profile.column_dimensions[column].width = 25

    analysis = workbook.create_sheet("Analysis")
    numbers = numeric_fields(rows)
    dimensions = dimension_fields(rows)
    metric = numbers[spec["number"] % len(numbers)]
    label_field = dimensions[0] if dimensions else fields[0]
    analysis.append([humanize(label_field), humanize(metric)])
    for row in rows:
        analysis.append([str(row[label_field]), float(row[metric])])
    if metric in FRACTION_PERCENT_FIELDS:
        for cell in analysis["B"][1:]:
            cell.number_format = "0.00%"
    elif metric in PERCENTAGE_POINT_FIELDS:
        for cell in analysis["B"][1:]:
            cell.number_format = '0.00"%"'
    style_excel_header(analysis[1], accent)
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = f"{humanize(metric)} by {humanize(label_field)}"
    chart.y_axis.title = humanize(metric)
    chart.x_axis.title = humanize(label_field)
    chart.add_data(Reference(analysis, min_col=2, min_row=1, max_row=len(rows) + 1), titles_from_data=True)
    chart.set_categories(Reference(analysis, min_col=1, min_row=2, max_row=len(rows) + 1))
    chart.height = 9
    chart.width = 17
    analysis.add_chart(chart, "D2")
    analysis.conditional_formatting.add(
        f"B2:B{len(rows) + 1}",
        ColorScaleRule(
            start_type="min",
            start_color="FCE8E6",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFF4CC",
            end_type="max",
            end_color="D9EAD3",
        ),
    )
    analysis.column_dimensions["A"].width = 28
    analysis.column_dimensions["B"].width = 20

    dictionary = workbook.create_sheet("Data Dictionary")
    dictionary.append(["Column", "Observed Type", "Description"])
    for column_definition in columns:
        dictionary.append(
            [
                column_definition["name"],
                column_definition["type"],
                column_definition["comment"],
            ]
        )
    style_excel_header(dictionary[1], accent)
    dictionary.column_dimensions["A"].width = 28
    dictionary.column_dimensions["B"].width = 18
    dictionary.column_dimensions["C"].width = 65

    decisions = workbook.create_sheet("Decision Notes")
    decisions.append(["Classification", "Observation or recommended follow-up"])
    for heading, detail in decision_notes(spec, rows, aggregate):
        decisions.append([heading, detail])
    style_excel_header(decisions[1], accent)
    decisions.column_dimensions["A"].width = 34
    decisions.column_dimensions["B"].width = 105
    for row in decisions.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    lineage = workbook.create_sheet("Lineage")
    lineage["A1"] = "Knowledge ingestion lineage"
    lineage["A1"].font = Font(size=18, bold=True, color=accent)
    for cell_range, heading, detail in (
        ("B3:D6", "Databricks", suffix(spec["source_table"])),
        ("F3:H6", "SharePoint", "Semiconductor Knowledge"),
        ("J3:L6", "Azure AI Search", "Chunked hybrid index"),
    ):
        lineage.merge_cells(cell_range)
        cell = lineage[cell_range.split(":")[0]]
        cell.value = f"{heading}\n{detail}"
        cell.fill = PatternFill("solid", fgColor="F3F7F8")
        cell.font = Font(size=14, bold=True, color=accent)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    lineage["E4"] = "->"
    lineage["I4"] = "->"
    for coordinate in ("E4", "I4"):
        lineage[coordinate].font = Font(size=24, bold=True, color=accent)
        lineage[coordinate].alignment = Alignment(horizontal="center")
    lineage.merge_cells("B9:L11")
    lineage["B9"] = narrative(spec["source_table"], rows, spec["profile_generated_at"])
    lineage["B9"].alignment = Alignment(wrap_text=True, vertical="top")
    for column in range(1, 13):
        lineage.column_dimensions[lineage.cell(1, column).column_letter].width = 13

    for worksheet in workbook.worksheets:
        worksheet.sheet_view.showGridLines = False
        worksheet.page_setup.fitToWidth = 1
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.oddFooter.center.text = f"{AUTHOR} | {spec['artifact_id']}"
    workbook.save(path)


def format_counts(document_count: int) -> dict[str, int]:
    base, remainder = divmod(document_count, len(FORMAT_ORDER))
    return {
        extension: base + (1 if index < remainder else 0)
        for index, extension in enumerate(FORMAT_ORDER)
    }


def build_specs(
    profile: dict, grouped: dict[str, list[dict]], document_count: int
) -> list[dict]:
    table_names = [table for table in profile["source_tables"] if table in grouped]
    if len(table_names) != 6:
        raise ValueError(f"Expected six sampled source tables, found {len(table_names)}")
    extensions = [
        extension
        for extension, count in format_counts(document_count).items()
        for _ in range(count)
    ]
    occurrence: Counter[tuple[str, str]] = Counter()
    specs = []
    for index, extension in enumerate(extensions, start=1):
        table_name = table_names[(index - 1) % len(table_names)]
        table_suffix = suffix(table_name)
        category, display_name, accent = TABLE_CONFIG[table_suffix]
        occurrence[(table_name, extension)] += 1
        variant = occurrence[(table_name, extension)]
        perspective = PERSPECTIVES[(index + variant - 2) % len(PERSPECTIVES)]
        title = f"{display_name}: {perspective} {variant}"
        artifact_id = f"SEM-{index:03d}"
        filename = f"{artifact_id.lower()}-{slug(title)}.{extension}"
        specs.append(
            {
                "number": index,
                "artifact_id": artifact_id,
                "title": title,
                "category": category,
                "perspective": perspective,
                "accent": accent,
                "extension": extension,
                "filename": filename,
                "relative_path": f"{slug(category)}/{filename}",
                "source_table": table_name,
                "source_system": SOURCE_SYSTEM,
                "profile_generated_at": profile["generated_at"],
            }
        )
    return specs


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate the 100-file Office corpus.")
    parser.add_argument("--profile", type=Path, default=Path("data/semiconductor_profile.json"))
    parser.add_argument("--output", type=Path, default=Path("corpus"))
    parser.add_argument("--document-count", type=int, default=100)
    args = parser.parse_args()
    if args.document_count < 1:
        parser.error("--document-count must be at least 1")

    profile = json.loads(args.profile.read_text(encoding="utf-8"))
    validate_profile(profile)
    grouped = group_rows(profile)
    specs = build_specs(profile, grouped, args.document_count)
    if args.output.exists():
        shutil.rmtree(args.output)
    args.output.mkdir(parents=True)

    manifest = []
    with tempfile.TemporaryDirectory(prefix="semiconductor-diagrams-") as temp_directory:
        temp_path = Path(temp_directory)
        for spec in specs:
            destination = args.output / spec["relative_path"]
            destination.parent.mkdir(parents=True, exist_ok=True)
            rows = sample_rows(grouped[spec["source_table"]], spec["number"])
            aggregate = aggregate_for_table(profile, spec["source_table"])
            columns = columns_for_table(profile, spec["source_table"])
            if spec["extension"] == "docx":
                diagram = temp_path / f"{spec['artifact_id']}.png"
                chart = temp_path / f"{spec['artifact_id']}-chart.png"
                create_lineage_image(diagram, spec["source_table"], spec["accent"])
                create_metric_chart_image(chart, rows, spec)
                create_docx(destination, spec, rows, aggregate, columns, diagram, chart)
            elif spec["extension"] == "pptx":
                create_pptx(destination, spec, rows, aggregate, columns)
            else:
                create_xlsx(destination, spec, rows, aggregate, columns)
            manifest.append(
                {
                    **{key: value for key, value in spec.items() if key not in {"accent", "number"}},
                    "row_count": len(rows),
                    "full_source_rows": aggregate.get("row_count"),
                    "source_column_count": len(columns),
                    "size_bytes": destination.stat().st_size,
                    "description": narrative(spec["source_table"], rows, spec["profile_generated_at"]),
                }
            )

    manifest_path = args.output / "manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                "documents": len(manifest),
                "formats": dict(Counter(item["extension"] for item in manifest)),
                "categories": dict(Counter(item["category"] for item in manifest)),
                "manifest": str(manifest_path),
            }
        )
    )


if __name__ == "__main__":
    main()