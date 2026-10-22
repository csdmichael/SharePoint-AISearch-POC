"""Text extraction for governance hashing and embedding."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from .policy import Policy
from .similarity import normalize_text

# Structural separators of extracted text: one line per paragraph and one cell per table column.
PARAGRAPH_SEPARATOR = "\n"
CELL_SEPARATOR = " | "


def extract_normalized_text(path: Path, policy: Policy) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(CELL_SEPARATOR.join(cell.text for cell in row.cells))
    return normalize_text(PARAGRAPH_SEPARATOR.join(parts), policy.whitespace_pattern())
