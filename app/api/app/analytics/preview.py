"""Structured, renderable previews of office and PDF documents for the in-page document viewer.

Unlike analytics/text.py (flat text for duplicate/conflict detection), these extractors keep enough
structure - pages, paragraph style, table rows, worksheet grids, slide titles, bullets, and embedded
diagrams/images - for the UI to render a rich multi-page preview directly in the browser.
"""

from __future__ import annotations

import base64
from io import BytesIO
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader

MAX_PARAGRAPHS = 800
MAX_TABLE_ROWS = 200
MAX_SHEET_ROWS = 300
MAX_SHEET_COLUMNS = 40
MAX_SHEETS = 25
MAX_SLIDES = 100
MAX_IMAGES_PER_DOC = 40


def _to_data_url(blob: bytes, content_type: str = "image/png") -> str:
    encoded = base64.b64encode(blob).decode("ascii")
    return f"data:{content_type};base64,{encoded}"


def _paragraph_has_page_break(paragraph: DocxParagraph) -> bool:
    if paragraph.paragraph_format and paragraph.paragraph_format.page_break_before:
        return True
    for br in paragraph._element.findall(".//" + qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return True
    return False


def _paragraph_images(paragraph: DocxParagraph, document: Document, remaining: int) -> list[str]:
    if remaining <= 0:
        return []
    urls: list[str] = []
    for blip in paragraph._element.findall(".//" + qn("a:blip")):
        if len(urls) >= remaining:
            break
        rel_id = blip.get(qn("r:embed"))
        if not rel_id or rel_id not in document.part.related_parts:
            continue
        part = document.part.related_parts[rel_id]
        content_type = getattr(part, "content_type", None) or "image/png"
        urls.append(_to_data_url(part.blob, content_type))
    return urls


def preview_docx(content: bytes) -> dict[str, Any]:
    document = Document(BytesIO(content))
    blocks: list[dict[str, Any]] = []
    tables: list[list[list[str]]] = []
    raw_pages: list[list[dict[str, Any]]] = [[]]
    image_count = 0

    for item in document.iter_inner_content():
        if isinstance(item, DocxParagraph):
            if _paragraph_has_page_break(item) and raw_pages[-1]:
                raw_pages.append([])
            text = item.text.strip()
            if text and len(blocks) < MAX_PARAGRAPHS:
                style_name = (item.style.name if item.style else "") or ""
                lower_style = style_name.lower()
                is_heading = lower_style.startswith("heading") or lower_style == "title"
                kind = "heading" if is_heading else "paragraph"
                entry = {"kind": kind, "text": text, "style": style_name}
                blocks.append(entry)
                raw_pages[-1].append(entry)
            for data_url in _paragraph_images(item, document, MAX_IMAGES_PER_DOC - image_count):
                image_count += 1
                raw_pages[-1].append({"kind": "image", "image": data_url})
            # Check trailing page break inside paragraph runs
            for br in item._element.findall(".//" + qn("w:br")):
                if br.get(qn("w:type")) == "page" and raw_pages[-1]:
                    raw_pages.append([])
                    break
        elif isinstance(item, DocxTable):
            rows = [[cell.text.strip() for cell in row.cells] for row in item.rows[:MAX_TABLE_ROWS]]
            tables.append(rows)
            raw_pages[-1].append({"kind": "table", "table": rows})

    non_empty_pages = [page for page in raw_pages if page]
    if not non_empty_pages and (blocks or tables):
        combined: list[dict[str, Any]] = list(blocks) + [{"kind": "table", "table": t} for t in tables]
        non_empty_pages = [combined]

    pages: list[dict[str, Any]] = []
    for idx, page_blocks in enumerate(non_empty_pages, start=1):
        heading = next((b.get("text") for b in page_blocks if b.get("kind") == "heading" and b.get("text")), None)
        pages.append({
            "number": idx,
            "title": heading or f"Page {idx}",
            "blocks": page_blocks,
        })

    return {
        "kind": "docx",
        "pageCount": len(pages),
        "pages": pages,
        "blocks": blocks,
        "tables": tables,
    }


def _shape_text_blocks(shapes, blocks: list[dict[str, Any]], image_budget: list[int]) -> None:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _shape_text_blocks(shape.shapes, blocks, image_budget)
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and image_budget[0] > 0:
            try:
                image_budget[0] -= 1
                content_type = getattr(shape.image, "content_type", None) or "image/png"
                blocks.append({"kind": "image", "image": _to_data_url(shape.image.blob, content_type)})
            except Exception:
                pass
            continue
        if shape.has_text_frame and shape.text_frame.text.strip():
            lines = [line.text.strip() for line in shape.text_frame.paragraphs if line.text.strip()]
            if lines:
                blocks.append({"lines": lines})
        if getattr(shape, "has_table", False) and shape.has_table:
            rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
            blocks.append({"table": rows})


def preview_pptx(content: bytes) -> dict[str, Any]:
    presentation = Presentation(BytesIO(content))
    slides: list[dict[str, Any]] = []
    image_budget = [MAX_IMAGES_PER_DOC]
    for index, slide in enumerate(presentation.slides, start=1):
        if index > MAX_SLIDES:
            break
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        blocks: list[dict[str, Any]] = []
        for shape in slide.shapes:
            if shape == slide.shapes.title:
                continue
            _shape_text_blocks([shape], blocks, image_budget)
        slides.append({"number": index, "title": title or f"Slide {index}", "blocks": blocks})
    return {"kind": "pptx", "pageCount": len(slides), "slides": slides}


def preview_xlsx(content: bytes) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(content), data_only=True)
    try:
        sheets: list[dict[str, Any]] = []
        image_count = 0
        for sheet in workbook.worksheets[:MAX_SHEETS]:
            rows: list[list[str]] = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= MAX_SHEET_ROWS:
                    break
                values = ["" if value is None else str(value) for value in row[:MAX_SHEET_COLUMNS]]
                if any(value.strip() for value in values):
                    rows.append(values)
            images: list[str] = []
            for img in getattr(sheet, "_images", []):
                if image_count >= MAX_IMAGES_PER_DOC:
                    break
                try:
                    raw = img._data()
                    if raw:
                        image_count += 1
                        images.append(_to_data_url(raw, "image/png"))
                except Exception:
                    pass
            sheet_entry: dict[str, Any] = {"name": sheet.title, "rows": rows}
            if images:
                sheet_entry["images"] = images
            sheets.append(sheet_entry)
        return {"kind": "xlsx", "pageCount": len(sheets), "sheets": sheets}
    finally:
        workbook.close()


def preview_pdf(content: bytes) -> dict[str, Any]:
    reader = PdfReader(BytesIO(content))
    pages: list[dict[str, Any]] = []
    image_count = 0
    for index, page in enumerate(reader.pages[:MAX_SLIDES], start=1):
        raw_text = (page.extract_text() or "").strip()
        lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
        title = lines[0] if lines else f"Page {index}"
        blocks: list[dict[str, Any]] = []
        for line_idx, line in enumerate(lines):
            if line_idx == 0:
                blocks.append({"kind": "heading", "text": line})
            elif "|" in line and len(line.split("|")) >= 2:
                cells = [c.strip() for c in line.split("|") if c.strip()]
                if blocks and blocks[-1].get("table") is not None:
                    blocks[-1]["table"].append(cells)
                else:
                    blocks.append({"kind": "table", "table": [cells]})
            else:
                blocks.append({"kind": "paragraph", "text": line})
        try:
            for img in getattr(page, "images", []):
                if image_count >= MAX_IMAGES_PER_DOC:
                    break
                if getattr(img, "data", None):
                    image_count += 1
                    blocks.append({"kind": "image", "image": _to_data_url(img.data, "image/png")})
        except Exception:
            pass
        pages.append({"number": index, "title": title, "blocks": blocks})
    return {"kind": "pdf", "pageCount": len(pages), "pages": pages}


PREVIEW_EXTRACTORS = {
    ".docx": preview_docx,
    ".pptx": preview_pptx,
    ".xlsx": preview_xlsx,
    ".pdf": preview_pdf,
}


def build_preview(extension: str, content: bytes) -> dict[str, Any]:
    extractor = PREVIEW_EXTRACTORS.get(extension.lower())
    if extractor is None:
        raise ValueError(f"no preview extractor for {extension}")
    return extractor(content)

