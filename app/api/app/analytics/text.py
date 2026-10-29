"""Text extraction and normalization for office docs and transcript text content.

Normalization, hashing and the numeric fingerprint follow the shared policy (similarity.* in
governance-policy.json). The extractors read only the document content that Microsoft Graph returns.
"""

from __future__ import annotations

import hashlib
import re
from io import BytesIO
from typing import Callable

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def _docx_text(content: bytes) -> str:
    document = Document(BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _shape_text(shapes, parts: list[str]) -> None:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            _shape_text(shape.shapes, parts)
            continue
        if shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))


def _pptx_text(content: bytes) -> str:
    presentation = Presentation(BytesIO(content))
    parts: list[str] = []
    for slide in presentation.slides:
        _shape_text(slide.shapes, parts)
    return "\n".join(parts)


def _xlsx_text(content: bytes) -> str:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None and str(value).strip()]
                if values:
                    parts.append(" ".join(values))
        return "\n".join(parts)
    finally:
        workbook.close()


def _plain_text(content: bytes) -> str:
    return content.decode("utf-8", errors="replace")


EXTRACTORS: dict[str, Callable[[bytes], str]] = {
    ".docx": _docx_text,
    ".pptx": _pptx_text,
    ".xlsx": _xlsx_text,
    ".txt": _plain_text,
    ".vtt": _plain_text,
}


def extract_text(extension: str, content: bytes) -> str:
    extractor = EXTRACTORS.get(extension.lower())
    if extractor is None:
        raise ValueError(f"no text extractor for {extension}")
    return extractor(content)


def normalize_text(text: str, whitespace_pattern: str) -> str:
    return re.sub(whitespace_pattern, " ", text).strip()


def content_digest(normalized: str, algorithm: str) -> str:
    return hashlib.new(algorithm, normalized.encode("utf-8")).hexdigest()


def numeric_fingerprint(normalized: str, number_pattern: str) -> tuple[str, ...]:
    return tuple(sorted(re.findall(number_pattern, normalized)))
