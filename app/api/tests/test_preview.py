"""Structured previews for the in-page document viewer."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from app.analytics.preview import build_preview


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Leave policy", level=1)
    document.add_paragraph("Primary caregivers receive 12 weeks.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Plan"
    table.rows[0].cells[1].text = "Weeks"
    table.rows[1].cells[0].text = "Primary"
    table.rows[1].cells[1].text = "12"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Plans"
    sheet.append(["Plan", "Premium"])
    sheet.append(["PPO", 145])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pptx_bytes() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Cohort overview"
    slide.placeholders[1].text_frame.text = "Eighteen managers"
    buffer = BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def test_docx_preview_separates_headings_paragraphs_and_tables() -> None:
    preview = build_preview(".docx", _docx_bytes())
    assert preview["kind"] == "docx"
    kinds = [block["kind"] for block in preview["blocks"]]
    assert kinds == ["heading", "paragraph"]
    assert preview["blocks"][0]["text"] == "Leave policy"
    assert preview["tables"] == [[["Plan", "Weeks"], ["Primary", "12"]]]


def test_xlsx_preview_returns_sheet_rows_as_text() -> None:
    preview = build_preview(".xlsx", _xlsx_bytes())
    assert preview["kind"] == "xlsx"
    assert preview["sheets"] == [{"name": "Plans", "rows": [["Plan", "Premium"], ["PPO", "145"]]}]


def test_pptx_preview_returns_titled_slides() -> None:
    preview = build_preview(".pptx", _pptx_bytes())
    assert preview["kind"] == "pptx"
    assert preview["slides"][0]["title"] == "Cohort overview"
    assert preview["slides"][0]["blocks"][0]["lines"] == ["Eighteen managers"]


def test_unsupported_extension_is_rejected() -> None:
    with pytest.raises(ValueError):
        build_preview(".mp4", b"")
