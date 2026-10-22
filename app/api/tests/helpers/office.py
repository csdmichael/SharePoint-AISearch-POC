"""Test-only builders for DOCX, PPTX and XLSX bytes from fixture content specs."""

from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches


def build_document(spec: dict[str, Any]) -> bytes:
    kind = spec["kind"]
    buffer = BytesIO()
    if kind == "docx":
        document = Document()
        for paragraph in spec["paragraphs"]:
            document.add_paragraph(paragraph)
        document.save(buffer)
    elif kind == "pptx":
        presentation = Presentation()
        blank = presentation.slide_layouts[6]
        for texts in spec["slides"]:
            slide = presentation.slides.add_slide(blank)
            for position, text in enumerate(texts):
                box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5 + position), Inches(8), Inches(0.8))
                box.text_frame.text = text
        presentation.save(buffer)
    elif kind == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        for row in spec["rows"]:
            sheet.append(row)
        workbook.save(buffer)
    else:
        raise ValueError(f"unsupported fixture kind {kind}")
    return buffer.getvalue()
